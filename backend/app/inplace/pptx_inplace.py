"""PPTX in-place bionic transformation.

Opens the user's original .pptx and walks every text run (`<a:r>`)
inside every shape and table cell of every slide. Runs are split into
bionic prefix-bold + suffix-normal pieces while preserving the run's
properties (`<a:rPr>`), so colors, fonts, sizes, hyperlinks, languages,
underline, italic, etc. are all kept intact.

What stays intact:
- Slide layouts, masters, themes (untouched)
- Images, shapes, smart-art, charts (untouched)
- Animations and transitions (untouched)
- Tables and their cells' formatting (preserved)
- Notes are processed (bionic-styled)
- Comments are untouched
"""
from __future__ import annotations

import io
from copy import deepcopy
from pathlib import Path

from lxml import etree
from pptx import Presentation

from ..models import BionicSettings
from ..transformer import WORD_RE, _prefix_length

_NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_TAG_R = f"{{{_NS_A}}}r"
_TAG_T = f"{{{_NS_A}}}t"
_TAG_RPR = f"{{{_NS_A}}}rPr"

_SAFE_RUN_CHILDREN = {_TAG_RPR, _TAG_T}


def export_inplace(
    data: bytes,
    settings: BionicSettings,
    filename: str = "presentation.pptx",
) -> tuple[bytes, str, str]:
    prs = Presentation(io.BytesIO(data))
    if settings.enabled:
        for slide in prs.slides:
            _process_shapes(slide.shapes, settings)
            if slide.has_notes_slide:
                _process_shapes(slide.notes_slide.shapes, settings)

    buf = io.BytesIO()
    prs.save(buf)
    stem = Path(filename).stem or "presentation"
    return (
        buf.getvalue(),
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        f"{stem}.bionic.pptx",
    )


def _process_shapes(shapes, settings: BionicSettings) -> None:
    for shape in shapes:
        try:
            if shape.shape_type == 6:  # MSO_SHAPE_TYPE.GROUP
                _process_shapes(shape.shapes, settings)
                continue
        except Exception:
            pass
        if shape.has_text_frame:
            _process_text_frame(shape.text_frame, settings)
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    _process_text_frame(cell.text_frame, settings)


def _process_text_frame(tf, settings: BionicSettings) -> None:
    for paragraph in tf.paragraphs:
        # Iterate over the underlying XML <a:r> children to support split.
        p_elem = paragraph._p
        runs = [c for c in p_elem if c.tag == _TAG_R]
        for r_elem in runs:
            _split_run(r_elem, settings)


def _split_run(r_elem, settings: BionicSettings) -> None:
    for child in r_elem:
        if child.tag not in _SAFE_RUN_CHILDREN:
            return

    text_elems = [c for c in r_elem if c.tag == _TAG_T]
    if not text_elems:
        return
    text = "".join((t.text or "") for t in text_elems)
    if not text.strip():
        return

    chunks = _compute_chunks(text, settings)
    if len(chunks) <= 1 or all(not is_prefix for _, is_prefix in chunks):
        return

    parent = r_elem.getparent()
    insert_at = list(parent).index(r_elem)

    new_runs = []
    for chunk_text, is_prefix in chunks:
        if not chunk_text:
            continue
        new_r = deepcopy(r_elem)
        for old_t in [c for c in new_r if c.tag == _TAG_T]:
            new_r.remove(old_t)
        if is_prefix:
            _set_bold(new_r)
        t = etree.SubElement(new_r, _TAG_T)
        t.text = chunk_text
        new_runs.append(new_r)

    for i, new_r in enumerate(new_runs):
        parent.insert(insert_at + i, new_r)
    parent.remove(r_elem)


def _set_bold(r_elem) -> None:
    rPr = r_elem.find(_TAG_RPR)
    if rPr is None:
        rPr = etree.Element(_TAG_RPR)
        r_elem.insert(0, rPr)
    rPr.set("b", "1")


def _compute_chunks(text: str, settings: BionicSettings) -> list[tuple[str, bool]]:
    chunks: list[tuple[str, bool]] = []
    last_end = 0
    for match in WORD_RE.finditer(text):
        start, end = match.span()
        if start > last_end:
            chunks.append((text[last_end:start], False))
        word = match.group(0)
        plen = _prefix_length(word, settings)
        if plen > 0:
            chunks.append((word[:plen], True))
            if plen < len(word):
                chunks.append((word[plen:], False))
        else:
            chunks.append((word, False))
        last_end = end
    if last_end < len(text):
        chunks.append((text[last_end:], False))
    return chunks
