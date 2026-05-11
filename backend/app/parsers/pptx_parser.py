"""PPTX parser — flattens slides/shapes/tables into DocumentModel blocks
for in-app reading. The user's original PPTX is never modified by this parser;
it's purely read-only.

For fidelity-preserving EXPORT use `app/inplace/pptx_inplace.py`.
"""
from __future__ import annotations

import io

from pptx import Presentation

from ..models import Block


def parse(data: bytes) -> tuple[list[Block], list[str]]:
    warnings: list[str] = []
    blocks: list[Block] = []
    try:
        prs = Presentation(io.BytesIO(data))
    except Exception as exc:
        warnings.append(f"Failed to open PPTX: {exc}")
        return blocks, warnings

    for slide_idx, slide in enumerate(prs.slides, start=1):
        blocks.append(Block(type="heading", text=f"Diapositive {slide_idx}", level=2))
        for shape in slide.shapes:
            _flatten_shape(shape, blocks)
        blocks.append(Block(type="spacer"))

    return blocks, warnings


def _flatten_shape(shape, blocks: list[Block]) -> None:
    if getattr(shape, "shape_type", None) == 6:  # GROUP
        for child in shape.shapes:
            _flatten_shape(child, blocks)
        return
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        for paragraph in shape.text_frame.paragraphs:
            text = "".join(run.text for run in paragraph.runs).strip()
            if text:
                blocks.append(Block(type="paragraph", text=text))
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = []
            for cell in row.cells:
                txt = "".join(
                    "".join(r.text for r in p.runs)
                    for p in cell.text_frame.paragraphs
                ).strip()
                cells.append(txt)
            row_text = " | ".join(c for c in cells if c)
            if row_text:
                blocks.append(Block(type="table_row", text=row_text))
