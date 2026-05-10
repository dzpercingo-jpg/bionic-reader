"""DOCX exporter — produces a NEW .docx with bionic prefixes applied as bold runs.

We do NOT modify the user's original file. We build a fresh document from the
parsed structure (headings, paragraphs, lists, blockquotes, code, table rows)
and split each word into two runs: a bold prefix and a normal suffix.
"""
from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt, RGBColor

from ..models import BionicSettings, DocumentModel
from ..transformer import WORD_RE, _prefix_length


def export(doc: DocumentModel, settings: BionicSettings, title: str | None = None) -> tuple[bytes, str, str]:
    out = Document()

    style = out.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    _ensure_quote_style(out)

    if title:
        out.add_heading(title, level=0)

    for block in doc.blocks:
        if block.type == "heading":
            level = max(1, min(block.level or 2, 6))
            p = out.add_heading(level=level)
            _add_bionic_runs(p, block.text or "", settings)
        elif block.type == "list_item":
            style_name = "List Number" if block.ordered else "List Bullet"
            try:
                p = out.add_paragraph(style=style_name)
            except KeyError:
                p = out.add_paragraph()
            _add_bionic_runs(p, block.text or "", settings)
        elif block.type == "blockquote":
            try:
                p = out.add_paragraph(style="IntenseQuote")
            except KeyError:
                p = out.add_paragraph()
            _add_bionic_runs(p, block.text or "", settings)
        elif block.type == "code":
            p = out.add_paragraph()
            run = p.add_run(block.text or "")
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        elif block.type == "table_row":
            p = out.add_paragraph()
            _add_bionic_runs(p, block.text or "", settings)
        elif block.type == "spacer":
            out.add_paragraph()
        else:
            p = out.add_paragraph()
            _add_bionic_runs(p, block.text or "", settings)

    buf = io.BytesIO()
    out.save(buf)
    stem = Path(doc.filename).stem or "document"
    return (
        buf.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        f"{stem}.bionic.docx",
    )


def _ensure_quote_style(d: Document) -> None:
    if "IntenseQuote" in [s.name for s in d.styles]:
        return
    try:
        d.styles.add_style("IntenseQuote", WD_STYLE_TYPE.PARAGRAPH)
    except Exception:
        pass


def _add_bionic_runs(paragraph, text: str, settings: BionicSettings) -> None:
    if not text:
        return
    if not settings.enabled:
        paragraph.add_run(text)
        return

    last_end = 0
    color = _hex_to_rgb(settings.prefix_color) if settings.use_color_instead_of_bold else None

    for match in WORD_RE.finditer(text):
        start, end = match.span()
        if start > last_end:
            paragraph.add_run(text[last_end:start])
        word = match.group(0)
        prefix_len = _prefix_length(word, settings)
        if prefix_len > 0:
            prefix = word[:prefix_len]
            suffix = word[prefix_len:]
            run = paragraph.add_run(prefix)
            run.bold = True
            if color:
                run.font.color.rgb = color
            if suffix:
                paragraph.add_run(suffix)
        else:
            paragraph.add_run(word)
        last_end = end
    if last_end < len(text):
        paragraph.add_run(text[last_end:])


def _hex_to_rgb(hex_color: str):
    h = hex_color.lstrip("#")
    if len(h) != 6:
        return RGBColor(0x11, 0x11, 0x11)
    try:
        r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
    except ValueError:
        return RGBColor(0x11, 0x11, 0x11)
    return RGBColor(r, g, b)
